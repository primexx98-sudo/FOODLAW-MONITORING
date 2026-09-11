"""사용자 로컬 PC의 표시(라벨링) 실무자료 폴더에서 텍스트를 추출해
data/practice_docs.json으로 저장 — embed_statutes.py가 이 파일을 읽어 기존 AI
의미검색 임베딩 파이프라인에 "실무자료" 출처로 합류시킨다.

**로컬 전용, 1회성 수동 실행 스크립트다** — 원본 파일이 GitHub Actions가 접근할 수
없는 사용자 PC 데스크톱에 있어서 collect_statutes.py처럼 매주 자동 실행할 수 없다.
새 자료를 추가하고 싶을 때마다 이 스크립트를 로컬에서 다시 실행하고, 결과
data/practice_docs.json을 커밋하는 방식으로 운영한다.

**공개 저장소 경고**: FOODLAW-MONITORING·food-monitor-hub 둘 다 Public GitHub
저장소이고 배포되는 사이트도 완전 공개다. 원본 폴더엔 실제 서명된 계약서·행정처분
공문·거래처 실명이 든 파일이 섞여 있어(2026-09-11 사용자와 함께 파일명 전수 검토),
EXCLUDE_PATHS에 그 항목들을 명시적으로 배제해뒀다 — 새 파일을 이 소스 폴더에
추가할 때마다 계약서/허가공문/실명 문서가 섞여 있는지 다시 확인하고 필요하면
EXCLUDE_PATHS를 갱신할 것. 또한 원본 폴더 경로 자체에 특정 인물명이 포함돼 있어
(Desktop\\<이름>\\표시\\...), **파일명만** 저장하고 원본 절대경로는 절대 저장하지
않는다(doc_id도 파일명 기반 해시로 생성, 개인 식별 정보 노출 방지).
"""

import hashlib
import json
import os
import re

import fitz  # pymupdf
import openpyxl
from pptx import Presentation

SOURCE_DIRS = [
    r"C:\Users\thefuture_brand\Desktop\이우찬\표시\★정보들",
    r"C:\Users\thefuture_brand\Desktop\이우찬\표시\공전",
]

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "practice_docs.json")

# 2026-09-11 사용자와 함께 파일명 전수 검토해 확정한 배제 목록 — 사유는 위 모듈
# docstring 참고. 폴더는 뒤에 os.sep를 붙여 하위 전체를 배제(부분 문자열로 다른
# 폴더가 우연히 걸리지 않도록).
EXCLUDE_PATHS = {
    # --- "개정된 사항" 배제(사용자 원 지시, 법령자료 탭의 diff/타임라인과 중복) ---
    "개정사항" + os.sep,
    "1. 식품등의 표시기준 일부개정고시(안).pdf",
    "26.01.01부터 적용되는 법령사항 - 타부서공유251023.pdf",
    "당류" + os.sep + "당알코올" + os.sep + "식품+등의+표시ㆍ광고에+관한+법률+시행규칙+일부개정령안(입법예고).pdf",
    "부자재" + os.sep + "2-2.'21년 분리배출 표시 지침 개정사항_공제조합교육자료.pdf",
    "식품 등의 표시ㆍ광고에 관한 법률 시행규칙_개정문개정이유.hwp",
    "일반식품" + os.sep + "과채가공품" + os.sep + "(1) 제1~제5_개정.hwpx",
    # --- 🔴 실제 계약서/허가공문/행정처분/브랜드 CI(거래처 실명 명시) ---
    "상표권 관련" + os.sep,
    "연구개발" + os.sep + "심의자료" + os.sep,
    "인쇄업체" + os.sep + "미래" + os.sep + "Signed agreement Korea Special Ink Ind. Co. 2024-1.pdf",
    "질의&유권해석" + os.sep + "[문산공장] 식품표시광고법 위반업체 행정처분통지 공문(변하게 하는 가차환).pdf",
    "로고관련" + os.sep + "삼성제약CI" + os.sep,
    "로고관련" + os.sep + "한미CI" + os.sep,
    "한미양행 패키지 ci.ai",
    "1. 250331 종근당 매뉴얼.ai",
    "자동포장" + os.sep + "S-8920A   한미양행 루테인 밀크씨슬 프리미엄(자동포장),,,,,,,,,,,,,,,,,날개 싸이즈 5mm 늘림.ai",
    # --- 🟡 애매 — 거래처명/개인명 포함 스펙시트·현황표(보수적으로 함께 배제) ---
    "240516 보령 루테인지아잔틴 미니 표시사항 체크리스트.xlsx",
    "엘케이에스 CAP 색상 List.xlsx",
    "부자재" + os.sep + "LKS 색상참고용 리스트.pdf",
    "PTP Change Part Tool Set 보유 현황표_HM.xlsx",
    "문산, 선유 공장 유형별 중요관리점(CCP).pdf",
    "표시업무 세미나_정은아.pptx",
    "PTP 도면" + os.sep,
    # --- 2026-09-11 2차 발견 — 파일명만으론 안 드러났지만 본문에 거래처 실명·담당자
    # 이메일·거래처 연락처가 포함된 것으로 확인(company-name/PII 스캔으로 발견) ---
    "QA 업무리스트.xlsx",
    "병 재질 리스트 양식(최종본).xlsx",
    "부자재 재질별 시험성적서 리스트, 부자재 및 원료업체 연락처.xlsx",
    "사이즈정리_.xlsx",
    "영양성분 계산기(건식)-1.xlsx",
    "영양성분계산식(일식)-1.xlsx",
}

# 이번 1차는 텍스트 파일부터(사용자 결정) — 이미지(jpg/png/tif)·디자인파일(ai/eps/psd)·
# 구형 .hwp(파서 없음)·기타(db/xls)는 스킵. 다음 라운드에서 OCR을 붙이면 이 집합을 넓히면 됨.
SUPPORTED_EXTS = {".pdf", ".xlsx", ".pptx", ".txt"}

_WS_RE = re.compile(r"\s")
# 2026-09-11 발견 — 파일명 검토만으론 못 걸러낸 개인 연락처(부자재 업체 담당자
# 전화번호·이메일 수십 건)가 "부자재 재질별 시험성적서 리스트..." 안에서 나옴 —
# EXCLUDE_PATHS로 그 문서 자체는 뺐지만, 앞으로 또 비슷한 게 다른 파일에 섞여 있을
# 위험에 대비해 **모든 추출 문서에 공통으로** 전화번호·이메일 패턴을 지운다(이중 방어).
_PHONE_RE = re.compile(r"01[016789][-.\s]?\d{3,4}[-.\s]?\d{4}")
_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")


def redact_pii(text: str) -> str:
    text = _EMAIL_RE.sub("[이메일 삭제]", text)
    text = _PHONE_RE.sub("[전화번호 삭제]", text)
    return text


def is_excluded(rel_path: str) -> bool:
    return any(rel_path == p or rel_path.startswith(p) for p in EXCLUDE_PATHS)


def extract_pdf_text(path: str) -> str:
    doc = fitz.open(path)
    return "\n".join(page.get_text() for page in doc)


def extract_xlsx_text(path: str) -> str:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_pptx_text(path: str) -> str:
    prs = Presentation(path)
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_txt_text(path: str) -> str:
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


EXTRACTORS = {
    ".pdf": extract_pdf_text,
    ".xlsx": extract_xlsx_text,
    ".pptx": extract_pptx_text,
    ".txt": extract_txt_text,
}


def looks_glued(text: str) -> bool:
    """한글 PDF 중 일부는 단어 사이 공백이 아예 없이 추출됨(2026-09-11 식품첨가물
    공전 PDF에서 실측: "적색리트머스지를쓴다") — 이런 파일은 검색 색인에 넣어봤자
    읽을 수 없는 텍스트라 신뢰도를 낮추므로 공백 비율로 걸러낸다. 정상적인 한국어
    문서는 공백이 전체 글자의 대략 10~20%를 차지한다(실측 hwpx 샘플 기준) — 그보다
    현저히 낮으면(5% 미만) 글자가 붙어서 나온 것으로 간주."""
    if len(text) < 200:
        return False
    ws = len(_WS_RE.findall(text))
    return (ws / len(text)) < 0.05


def make_doc_id(filename: str) -> str:
    return "practice_" + hashlib.md5(filename.encode("utf-8")).hexdigest()[:12]


def clean_title(filename: str) -> str:
    name = os.path.splitext(filename)[0]
    name = name.replace("+", " ")
    name = re.sub(r"\s+", " ", name).strip()
    name = name.lstrip("★").strip()
    name = re.sub(r"^\d+\.\s+", "", name)  # "1. 제목"류 번호매김 접두사만 제거(연도 등 의미있는 숫자는 보존)
    return name.strip()


def collect() -> None:
    docs = []
    seen_hashes = set()
    skipped_excluded = skipped_format = skipped_glued = skipped_error = skipped_dup = 0

    for source_dir in SOURCE_DIRS:
        if not os.path.isdir(source_dir):
            print(f"[실무자료] 소스 폴더 없음, 건너뜀: {source_dir}")
            continue
        for root, _, files in os.walk(source_dir):
            for filename in files:
                full_path = os.path.join(root, filename)
                rel_path = os.path.relpath(full_path, source_dir)
                if is_excluded(rel_path):
                    skipped_excluded += 1
                    continue
                ext = os.path.splitext(filename)[1].lower()
                if ext not in SUPPORTED_EXTS:
                    skipped_format += 1
                    continue
                try:
                    text = redact_pii(EXTRACTORS[ext](full_path).strip())
                except Exception as e:
                    print(f"[실무자료] 추출 실패({filename}): {e}")
                    skipped_error += 1
                    continue
                if len(text) < 50:  # 거의 빈 문서(도면 PDF 등 이미지 위주) — 검색 가치 없음
                    skipped_error += 1
                    continue
                if looks_glued(text):
                    print(f"[실무자료] 공백 없이 붙어서 추출됨(품질 낮음) — 제외: {filename}")
                    skipped_glued += 1
                    continue
                # 같은 문서가 여러 폴더에 복사돼 있는 경우(실측: 유기가공식품 질의응답자료집이
                # 3곳에 중복 존재)가 있어 — 내용 동일본은 첫 번째만 남기고 스킵.
                content_hash = hashlib.md5(text.encode("utf-8")).hexdigest()
                if content_hash in seen_hashes:
                    skipped_dup += 1
                    continue
                seen_hashes.add(content_hash)
                docs.append({
                    "id": make_doc_id(filename),
                    "name": clean_title(filename),
                    "text": text,
                })

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump({"docs": docs}, f, ensure_ascii=False, indent=2)

    total_chars = sum(len(d["text"]) for d in docs)
    print(f"\n[실무자료] 완료: {len(docs)}건 추출({total_chars:,}자) → {OUT_PATH}")
    print(f"[실무자료] 배제(제외목록) {skipped_excluded}건 · 미지원형식 {skipped_format}건 · "
          f"추출실패/빈문서 {skipped_error}건 · 저품질(공백없음) {skipped_glued}건 · 중복 {skipped_dup}건")


if __name__ == "__main__":
    collect()
